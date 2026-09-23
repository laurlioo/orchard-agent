"""报表聚合 + Excel 导出。"""
from datetime import date, timedelta
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.product import ProductCategory
from app.models.production_log import ProductionLog
from app.schemas.report import CategoryRow, ReportData
from app.services.wage_calculator import calc_wages


def _infer_report_type(start: date, end: date) -> str:
    span = (end - start).days + 1
    if span == 1:
        return "daily"
    if span == 7:
        return "weekly"
    if span >= 28 and span <= 31:
        return "monthly"
    return "custom"


def build_report(db: Session, start: date, end: date, orchard: str = "peach") -> ReportData:
    """聚合 [start, end] 区间内的产量 + 工资数据。"""
    # 各品类聚合
    cat_rows = (
        db.query(
            ProductCategory.id.label("category_id"),
            ProductCategory.name,
            ProductCategory.unit,
            ProductCategory.unit_price,
            ProductCategory.cost_per_unit,
            func.sum(ProductionLog.quantity).label("quantity"),
            func.sum(
                ProductionLog.quantity
                * func.coalesce(ProductionLog.unit_price, ProductCategory.unit_price)
            ).label("revenue"),
        )
        .join(ProductionLog, ProductionLog.category_id == ProductCategory.id)
        .filter(
            ProductionLog.date >= start,
            ProductionLog.date <= end,
            ProductionLog.orchard == orchard,
        )
        .group_by(ProductCategory.id)
        .all()
    )
    categories = []
    for r in cat_rows:
        quantity = float(r.quantity or 0)
        revenue = round(float(r.revenue or 0), 2)
        cost = round(quantity * float(r.cost_per_unit or 0), 2)
        gross_profit = round(revenue - cost, 2)
        unit_price = round(revenue / quantity, 2) if quantity else float(r.unit_price or 0)
        categories.append(
            CategoryRow(
                category_id=r.category_id,
                name=r.name,
                unit=r.unit or "斤",
                quantity=quantity,
                unit_price=unit_price,
                cost_per_unit=float(r.cost_per_unit or 0),
                revenue=revenue,
                cost=cost,
                gross_profit=gross_profit,
            )
        )

    workers = calc_wages(db, start, end, orchard=orchard)

    return ReportData(
        period_start=start,
        period_end=end,
        report_type=_infer_report_type(start, end),
        categories=categories,
        workers=workers,
        total_quantity=round(sum(c.quantity for c in categories), 2),
        total_revenue=round(sum(c.revenue for c in categories), 2),
        total_cost=round(sum(c.cost for c in categories), 2),
        total_gross_profit=round(sum(c.gross_profit for c in categories), 2),
        total_hours=round(sum(w.hours for w in workers), 2),
        total_overtime_hours=round(sum(w.overtime_hours for w in workers), 2),
        total_regular_wages=round(sum(w.regular_wage for w in workers), 2),
        total_overtime_wages=round(sum(w.overtime_wage for w in workers), 2),
        total_wages=round(sum(w.wage for w in workers), 2),
    )


def export_excel(report: ReportData) -> bytes:
    """生成 .xlsx 字节流，含汇总 sheet + 明细 sheet。"""
    wb = Workbook()

    # Sheet 1: 汇总
    ws = wb.active
    ws.title = "汇总"
    ws.append(
        [
            f"{report.report_type} 报表",
            f"{report.period_start} ~ {report.period_end}",
        ]
    )
    ws.append([])
    ws.append(["指标", "数值"])
    ws.append(["总产量", report.total_quantity])
    ws.append(["总销售额", report.total_revenue])
    ws.append(["总成本", report.total_cost])
    ws.append(["总毛利", report.total_gross_profit])
    ws.append(["正常工时", report.total_hours])
    ws.append(["加班工时", report.total_overtime_hours])
    ws.append(["正常工资", report.total_regular_wages])
    ws.append(["加班工资", report.total_overtime_wages])
    ws.append(["总工资", report.total_wages])
    if report.summary:
        ws.append([])
        ws.append(["AI 摘要", report.summary])

    # Sheet 2: 品类明细
    ws2 = wb.create_sheet("品类明细")
    ws2.append(
        [
            "品类",
            "单位",
            "数量",
            "单价",
            "单位成本",
            "销售额",
            "成本",
            "毛利",
        ]
    )
    for c in report.categories:
        ws2.append(
            [
                c.name,
                c.unit,
                c.quantity,
                c.unit_price,
                c.cost_per_unit,
                c.revenue,
                c.cost,
                c.gross_profit,
            ]
        )

    # Sheet 3: 工资明细
    ws3 = wb.create_sheet("工资明细")
    ws3.append(
        [
            "工人",
            "岗位",
            "正常工时",
            "加班工时",
            "时薪",
            "加班倍数",
            "正常工资",
            "加班工资",
            "总工资",
        ]
    )
    for w in report.workers:
        ws3.append(
            [
                w.name,
                w.role,
                w.hours,
                w.overtime_hours,
                w.hourly_rate,
                w.overtime_rate,
                w.regular_wage,
                w.overtime_wage,
                w.wage,
            ]
        )

    # 简单列宽自适应
    for sheet in [ws, ws2, ws3]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
